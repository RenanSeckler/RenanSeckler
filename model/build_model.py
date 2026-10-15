"""Gera model/Seckler_Modelo_Financeiro.xlsx (modelo editável com fórmulas).

Uso:  python3 model/build_model.py
Células azuis = premissas editáveis. Todo o resto é fórmula.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("Seckler_Modelo_Financeiro.xlsx")

INPUT = PatternFill("solid", fgColor="DDEBF7")
HEAD = PatternFill("solid", fgColor="1F3864")
SECTION = PatternFill("solid", fgColor="D9D9D9")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
USD = '"$"#,##0;[Red]-"$"#,##0'
PCT = "0.0%"
NUM = "#,##0"

CASES = ["Conservador", "Base", "Otimista"]
STRUCTS = ["A - Terceirizar", "B - Alugar", "C - Comprar"]

# (chave, rótulo, unidade, [Cons, Base, Otim], formato, fonte/status)
CASE_INPUTS = [
    ("jobs", "Mudanças por mês", "jobs", [8, 14, 22], NUM,
     "PENDENTE: histórico real da Seckler. FMCSA mostra 1 caminhão/1 motorista (localmovers.com)."),
    ("ticket", "Ticket médio por mudança", "US$", [850, 1100, 1300], USD,
     "moveBuddha Orlando: US$672-1.550 (1-2 quartos); FreightWaves: US$306-2.712. Validar com faturas."),
    ("cv_mov", "Custo variável da mudança", "% receita", [0.58, 0.55, 0.52], PCT,
     "PREMISSA: mão de obra, combustível, materiais, comissões. Substituir pela DRE real."),
    ("packouts", "Pack-outs (sinistros) por mês", "jobs", [0, 1, 3], NUM,
     "PREMISSA: zero no conservador até existir parceiro de restauração assinado."),
    ("ticket_po", "Ticket médio pack-out + devolução (sem armazenagem)", "US$", [2500, 3500, 4500], USD,
     "Referência única de fornecedor (Maryland): US$1.500-6.500 por residência. Baixa confiança."),
    ("cv_po", "Custo variável do pack-out", "% receita", [0.55, 0.50, 0.45], PCT,
     "PREMISSA: mão de obra intensiva, caixas, inventário, transporte."),
    ("vaults", "Vaults ocupados (média do período)", "vaults", [8, 25, 50], NUM,
     "PREMISSA: estoque médio de clientes em armazenagem (mudança + sinistro)."),
    ("price_v", "Preço por vault / mês", "US$", [90, 110, 130], USD,
     "PREMISSA a validar por cotação oculta (1st Class, Suddath, Cento, J&J Metro). "
     "Self-storage 10x10 climatizado WG: ~US$80-155/mês (RentCafe/Public Storage 2026)."),
    ("handling", "Manuseio + seguro de armazenagem por vault / mês", "US$", [10, 8, 7], USD,
     "PREMISSA."),
    ("dso", "Prazo médio de recebimento do pack-out (DSO)", "dias", [75, 60, 45], NUM,
     "PREMISSA: restauradoras/seguradoras pagam com atraso; validar nos contratos."),
    ("reserve", "Reserva de caixa", "meses de custo fixo", [3, 2, 2], NUM,
     "Política de prudência."),
    ("gov", "Receita governo (SOMENTE contratos assinados)", "US$/mês", [0, 0, 0], USD,
     "Regra: propostas e leads não entram. Atualizar apenas após adjudicação."),
]

STRUCT_INPUTS = [
    ("area", "Área do imóvel", "sf", [0, 3000, 6000], NUM,
     "B: listagens WG/Ocoee 2.200-5.255 sf. C: comparáveis Ocoee/Apopka 6.040-6.500 sf."),
    ("cap", "Capacidade de vaults", "vaults", [100000, 100, 200], NUM,
     "PREMISSA: ~60% da área útil, vault 7x5 ft, empilhado 2 alturas. A = capacidade do parceiro."),
    ("out_cost", "Custo terceirizado por vault / mês", "US$", [60, 0, 0], USD,
     "PREMISSA central do cenário A: cotar com armazéns de mudança/3PL. Ver aba Sensibilidade."),
    ("rent", "Aluguel base", "US$/sf/ano", [0, 17.5, 0], '"$"#,##0.00',
     "LoopNet/Showcase out/2026: 661 Garden Commerce Pky US$16,50-16,75; Business Park Blvd US$17-19,20; "
     "Maguire Rd Ext Ocoee US$17. Imposto sobre aluguel comercial revogado em 01/10/2025 (HB 7031)."),
    ("cam", "CAM/NNN (impostos, seguro, áreas comuns)", "US$/sf/ano", [0, 4, 0], '"$"#,##0.00',
     "PREMISSA: listagens não informam NNN. Confirmar com corretor."),
    ("utilities", "Utilidades + manutenção", "US$/mês", [0, 600, 900], USD, "PREMISSA."),
    ("wll", "Seguro warehouse legal liability", "US$/mês", [0, 300, 300], USD,
     "PREMISSA: cotar com corretor de seguros."),
    ("base_fixed", "Custo fixo base (seguros auto/carga/GL, software, marketing, admin)", "US$/mês",
     [3500, 3500, 3500], USD, "PENDENTE: substituir pelos custos fixos reais da Seckler."),
    ("staff", "Pessoal fixo adicional do armazém", "US$/mês", [0, 1500, 2500], USD,
     "PREMISSA: meio período (B) / integral parcial (C)."),
    ("sublease", "Receita de sublocação (área excedente)", "US$/mês", [0, 0, 3500], USD,
     "C: ~2.500 sf a ~US$17/sf/ano. SBA 504 exige ocupação própria (51% em imóvel existente — verificar SOP)."),
    ("price_sf", "Preço de compra", "US$/sf", [0, 0, 300], USD,
     "Pedidos: 300 Story Rd Ocoee ~US$385/sf; Clarcona Rd Apopka ~US$364/sf; média WG ~US$250/sf (agregador)."),
    ("equity", "Entrada (equity)", "% do preço", [0, 0, 0.10], PCT,
     "SBA 504 padrão 10%; 15% se empresa <2 anos ou imóvel de uso especial."),
    ("bank_pct", "Parcela banco (1ª hipoteca)", "% do preço", [0, 0, 0.50], PCT, "Estrutura 504 50/40/10."),
    ("bank_rate", "Juros banco", "% a.a.", [0, 0, 0.0725], '0.00%', "PREMISSA: cotar com bancos/CDC."),
    ("cdc_pct", "Parcela CDC/SBA (debênture)", "% do preço", [0, 0, 0.40], PCT, "Estrutura 504."),
    ("cdc_rate", "Taxa efetiva CDC (com encargos)", "% a.a.", [0, 0, 0.065], '0.00%',
     "PREMISSA: taxa da debênture é fixada no mês da venda; cotar com CDC (ex.: Florida First Capital)."),
    ("term", "Prazo do financiamento", "anos", [0, 0, 25], NUM, "504 imóvel: 10/20/25 anos."),
    ("closing", "Custos de fechamento", "% do preço", [0, 0, 0.04], PCT,
     "PREMISSA: avaliação, Phase I ESA, título, taxas SBA/CDC."),
    ("tax_ins", "IPTU + seguro do imóvel", "% do preço/ano", [0, 0, 0.025], PCT,
     "PREMISSA: Orange County ~1,6-1,9% + seguro patrimonial FL. Confirmar no Property Appraiser."),
    ("deposit", "Caução + 1º mês", "meses de aluguel", [0, 3, 0], NUM, "PREMISSA comum: 2 caução + 1º mês."),
    ("forklift", "Empilhadeira (usada)", "US$", [0, 18000, 18000], USD, "PREMISSA: cotar."),
    ("vault_qty", "Vaults comprados inicialmente", "un", [0, 50, 80], NUM, "Comprar em lotes conforme ocupação."),
    ("vault_cost", "Custo por vault", "US$", [0, 375, 375], USD, "PREMISSA: cotar fabricantes."),
    ("fitout", "Adequação, licenças, incêndio, segurança", "US$", [0, 12000, 25000], USD,
     "PREMISSA: depende de zoneamento e inspeção (Fire Marshal)."),
    ("equip", "Software de inventário, etiquetas, equipamentos", "US$", [4000, 6000, 6000], USD,
     "Inventário digital com fotos/código de barras é requisito prático de seguradoras."),
]


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def build_premissas(wb):
    ws = wb.active
    ws.title = "Premissas"
    ws["A1"] = "Premissas — células azuis são editáveis. Fonte/status na coluna F."
    ws["A1"].font = Font(bold=True, size=13)
    refs = {}

    ws.append([])
    ws.append(["Parâmetro (varia por CASO)", "Unidade", *CASES, "Fonte / status"])
    case_hdr = ws.max_row
    style_header(ws, case_hdr, 6)
    for key, label, unit, vals, fmt, src in CASE_INPUTS:
        ws.append([label, unit, *vals, src])
        r = ws.max_row
        for c in range(3, 6):
            ws.cell(row=r, column=c).fill = INPUT
            ws.cell(row=r, column=c).number_format = fmt
        refs[key] = ("case", r)

    ws.append([])
    ws.append(["Parâmetro (varia por ESTRUTURA)", "Unidade", *STRUCTS, "Fonte / status"])
    struct_hdr = ws.max_row
    style_header(ws, struct_hdr, 6)
    for key, label, unit, vals, fmt, src in STRUCT_INPUTS:
        ws.append([label, unit, *vals, src])
        r = ws.max_row
        for c in range(3, 6):
            ws.cell(row=r, column=c).fill = INPUT
            ws.cell(row=r, column=c).number_format = fmt
        refs[key] = ("struct", r)

    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18
    for col in "CDE":
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["F"].width = 110
    ws.freeze_panes = "C4"
    return refs, case_hdr, struct_hdr


def build_cenarios(wb, refs, case_hdr, struct_hdr):
    ws = wb.create_sheet("Cenarios")
    ws["A1"] = "Resultado mensal em regime (US$) — 3 estruturas x 3 casos. Não editar: tudo é fórmula."
    ws["A1"].font = Font(bold=True, size=13)
    combos = [(s, c) for s in STRUCTS for c in CASES]
    ws.cell(row=2, column=1, value="Estrutura")
    ws.cell(row=3, column=1, value="Caso")
    for i, (s, c) in enumerate(combos):
        ws.cell(row=2, column=2 + i, value=s)
        ws.cell(row=3, column=2 + i, value=c)
    style_header(ws, 2, 1 + len(combos))
    style_header(ws, 3, 1 + len(combos))

    def p(key, col):
        kind, r = refs[key]
        hdr = case_hdr if kind == "case" else struct_hdr
        sel = f"{col}$3" if kind == "case" else f"{col}$2"
        return (f"INDEX(Premissas!$C${r}:$E${r},MATCH({sel},Premissas!$C${hdr}:$E${hdr},0))")

    rows = {}
    lines = [
        ("sec", "RECEITA"),
        ("rev_mov", "Receita de mudanças", lambda c: f"={p('jobs', c)}*{p('ticket', c)}", USD),
        ("rev_po", "Receita de pack-out (sinistros)", lambda c: f"={p('packouts', c)}*{p('ticket_po', c)}", USD),
        ("vaults", "Vaults ocupados (limitado à capacidade)", lambda c: f"=MIN({p('vaults', c)},{p('cap', c)})", NUM),
        ("rev_st", "Receita de armazenagem", lambda c: f"={c}{{vaults}}*{p('price_v', c)}", USD),
        ("rev_gov", "Receita governo (contratos assinados)", lambda c: f"={p('gov', c)}", USD),
        ("rev_sub", "Receita de sublocação", lambda c: f"={p('sublease', c)}", USD),
        ("rev", "RECEITA TOTAL", lambda c: f"=SUM({c}{{rev_mov}},{c}{{rev_po}},{c}{{rev_st}},{c}{{rev_gov}},{c}{{rev_sub}})", USD),
        ("sec", "CUSTOS VARIÁVEIS"),
        ("cv_mov", "CV mudanças", lambda c: f"={c}{{rev_mov}}*{p('cv_mov', c)}", USD),
        ("cv_po", "CV pack-out", lambda c: f"={c}{{rev_po}}*{p('cv_po', c)}", USD),
        ("cv_st", "CV armazenagem (manuseio + custo terceirizado)", lambda c: f"={c}{{vaults}}*({p('handling', c)}+{p('out_cost', c)})", USD),
        ("cv_gov", "CV governo (premissa 60%)", lambda c: f"={c}{{rev_gov}}*0.6", USD),
        ("cv", "CUSTO VARIÁVEL TOTAL", lambda c: f"=SUM({c}{{cv_mov}},{c}{{cv_po}},{c}{{cv_st}},{c}{{cv_gov}})", USD),
        ("mc", "MARGEM DE CONTRIBUIÇÃO", lambda c: f"={c}{{rev}}-{c}{{cv}}", USD),
        ("mc_pct", "Margem de contribuição %", lambda c: f"=IF({c}{{rev}}>0,{c}{{mc}}/{c}{{rev}},0)", PCT),
        ("sec", "CUSTOS FIXOS (mensal)"),
        ("f_base", "Custo fixo base", lambda c: f"={p('base_fixed', c)}", USD),
        ("f_rent", "Aluguel + CAM", lambda c: f"={p('area', c)}*({p('rent', c)}+{p('cam', c)})/12", USD),
        ("f_util", "Utilidades + manutenção", lambda c: f"={p('utilities', c)}", USD),
        ("f_wll", "Seguro de armazém", lambda c: f"={p('wll', c)}", USD),
        ("f_staff", "Pessoal do armazém", lambda c: f"={p('staff', c)}", USD),
        ("prop_val", "Valor do imóvel (memória)", lambda c: f"={p('area', c)}*{p('price_sf', c)}", USD),
        ("f_bank", "Parcela banco", lambda c: f"=IF({p('term', c)}>0,PMT({p('bank_rate', c)}/12,{p('term', c)}*12,-{c}{{prop_val}}*{p('bank_pct', c)}),0)", USD),
        ("f_cdc", "Parcela CDC/SBA 504", lambda c: f"=IF({p('term', c)}>0,PMT({p('cdc_rate', c)}/12,{p('term', c)}*12,-{c}{{prop_val}}*{p('cdc_pct', c)}),0)", USD),
        ("f_tax", "IPTU + seguro do imóvel", lambda c: f"={c}{{prop_val}}*{p('tax_ins', c)}/12", USD),
        ("fixed", "CUSTO FIXO TOTAL", lambda c: f"=SUM({c}{{f_base}},{c}{{f_rent}},{c}{{f_util}},{c}{{f_wll}},{c}{{f_staff}},{c}{{f_bank}},{c}{{f_cdc}},{c}{{f_tax}})", USD),
        ("sec", "RESULTADO"),
        ("op", "Resultado operacional de caixa / mês", lambda c: f"={c}{{mc}}-{c}{{fixed}}", USD),
        ("op_y", "Resultado operacional de caixa / ano", lambda c: f"={c}{{op}}*12", USD),
        ("be_rev", "Ponto de equilíbrio (receita/mês)", lambda c: f"=IF({c}{{mc_pct}}>0,{c}{{fixed}}/{c}{{mc_pct}},0)", USD),
        ("be_v", "Vaults necessários p/ equilíbrio (demais receitas constantes)",
         lambda c: f"=MAX(0,({c}{{fixed}}-({c}{{mc}}-{c}{{vaults}}*({p('price_v', c)}-{p('handling', c)}-{p('out_cost', c)})))/({p('price_v', c)}-{p('handling', c)}-{p('out_cost', c)}))", NUM),
        ("be_ok", "Equilíbrio cabe na capacidade?", lambda c: f'=IF({c}{{be_v}}<={p("cap", c)},"SIM","NÃO")', None),
        ("occ", "Ocupação do armazém", lambda c: f"=IF({p('cap', c)}<100000,{c}{{vaults}}/{p('cap', c)},0)", PCT),
        ("sec", "INVESTIMENTO INICIAL E CAPITAL DE GIRO"),
        ("i_dep", "Caução + 1º mês", lambda c: f"={p('deposit', c)}*{c}{{f_rent}}", USD),
        ("i_eq", "Entrada do imóvel", lambda c: f"={c}{{prop_val}}*{p('equity', c)}", USD),
        ("i_close", "Custos de fechamento", lambda c: f"={c}{{prop_val}}*{p('closing', c)}", USD),
        ("i_fork", "Empilhadeira", lambda c: f"={p('forklift', c)}", USD),
        ("i_vault", "Vaults", lambda c: f"={p('vault_qty', c)}*{p('vault_cost', c)}", USD),
        ("i_fit", "Adequação / licenças / segurança", lambda c: f"={p('fitout', c)}", USD),
        ("i_eqp", "Software e equipamentos", lambda c: f"={p('equip', c)}", USD),
        ("inv", "INVESTIMENTO INICIAL TOTAL", lambda c: f"=SUM({c}{{i_dep}},{c}{{i_eq}},{c}{{i_close}},{c}{{i_fork}},{c}{{i_vault}},{c}{{i_fit}},{c}{{i_eqp}})", USD),
        ("wc", "Capital de giro (recebíveis pack-out + reserva)", lambda c: f"={c}{{rev_po}}*{p('dso', c)}/30+{p('reserve', c)}*{c}{{fixed}}", USD),
        ("cash", "NECESSIDADE TOTAL DE CAIXA", lambda c: f"={c}{{inv}}+{c}{{wc}}", USD),
        ("payback", "Payback do investimento (meses)", lambda c: f'=IF({c}{{op}}>0,{c}{{inv}}/{c}{{op}},"não recupera")', '0.0'),
    ]

    r = 4
    pending = []
    for item in lines:
        if item[0] == "sec":
            ws.cell(row=r, column=1, value=item[1]).font = BOLD
            for col in range(1, 2 + len(combos)):
                ws.cell(row=r, column=col).fill = SECTION
            r += 1
            continue
        key, label, fn, fmt = item
        rows[key] = r
        ws.cell(row=r, column=1, value=label)
        if label.isupper():
            ws.cell(row=r, column=1).font = BOLD
        pending.append((r, fn, fmt, label.isupper()))
        r += 1

    for row, fn, fmt, bold in pending:
        for i in range(len(combos)):
            col = get_column_letter(2 + i)
            formula = fn(col).replace("{{", "{").replace("}}", "}").format(**rows)
            cell = ws.cell(row=row, column=2 + i, value=formula)
            if fmt:
                cell.number_format = fmt
            if bold:
                cell.font = BOLD

    last = get_column_letter(1 + len(combos))
    red = PatternFill("solid", fgColor="F8CBAD")
    green = PatternFill("solid", fgColor="C6EFCE")
    rng = f"B{rows['op']}:{last}{rows['op']}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=red))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=green))

    ws.column_dimensions["A"].width = 58
    for i in range(len(combos)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 17
    ws.freeze_panes = "B4"
    return rows


def build_sensibilidade(wb, refs, struct_hdr):
    """Ponto de indiferença: quantos vaults estáveis justificam sair do cenário A."""
    ws = wb.create_sheet("Sensibilidade")
    ws["A1"] = "Quando vale a pena ter armazém próprio? (vaults ocupados estáveis em que B ou C empatam com A)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Lógica: armazém próprio troca custo variável por vault (terceirizado) por custo fixo. "
                "Indiferença = (Fixo extra da estrutura − receita de sublocação) ÷ custo terceirizado por vault.")
    ws.append([])
    ws.append(["Custo terceirizado por vault/mês (A)", "Vaults p/ B empatar com A", "Capacidade B",
               "Vaults p/ C empatar com A", "Capacidade C"])
    style_header(ws, 4, 5)
    c_rows = {k: refs[k][1] for k in ("base_fixed", "area", "rent", "cam", "utilities", "wll", "staff",
                                       "price_sf", "bank_pct", "bank_rate", "cdc_pct", "cdc_rate", "term",
                                       "tax_ins", "sublease", "cap")}

    def fixed(col):
        g = lambda k: f"Premissas!${col}${c_rows[k]}"
        val = f"({g('area')}*{g('price_sf')})"
        debt = (f"IF({g('term')}>0,PMT({g('bank_rate')}/12,{g('term')}*12,-{val}*{g('bank_pct')})"
                f"+PMT({g('cdc_rate')}/12,{g('term')}*12,-{val}*{g('cdc_pct')}),0)")
        return (f"({g('area')}*({g('rent')}+{g('cam')})/12+{g('utilities')}+{g('wll')}+{g('staff')}"
                f"+{debt}+{val}*{g('tax_ins')}/12-{g('sublease')})")

    for i, cost in enumerate([40, 50, 60, 70, 80, 90, 100, 120]):
        r = 5 + i
        ws.cell(row=r, column=1, value=cost).fill = INPUT
        ws.cell(row=r, column=1).number_format = USD
        ws.cell(row=r, column=2, value=f"=ROUNDUP({fixed('D')}/A{r},0)").number_format = NUM
        ws.cell(row=r, column=3, value=f"=Premissas!$D${c_rows['cap']}").number_format = NUM
        ws.cell(row=r, column=4, value=f"=ROUNDUP({fixed('E')}/A{r},0)").number_format = NUM
        ws.cell(row=r, column=5, value=f"=Premissas!$E${c_rows['cap']}").number_format = NUM
    ws.cell(row=14, column=1, value=("Leitura: se o número de vaults para empatar for maior que a capacidade, "
                                     "a estrutura só se paga por outros motivos (contratos que exigem instalação "
                                     "própria, capacidade de pack-out, patrimônio).")).alignment = Alignment(wrap_text=True)
    ws.merge_cells("A14:E16")
    for col, w in zip("ABCDE", (36, 26, 16, 26, 16)):
        ws.column_dimensions[col].width = w


def build_crm(wb):
    ws = wb.create_sheet("CRM")
    hdr = ["ID", "Pipeline", "Empresa", "Contato (nome/cargo)", "Telefone/canal público", "Fonte verificada",
           "Prioridade", "Etapa", "Responsável", "Último contato", "Próximo passo", "Data follow-up",
           "Valor potencial/mês (US$)", "Probabilidade", "Valor ponderado", "Notas"]
    ws.append(hdr)
    style_header(ws, 1, len(hdr))
    stages = ("0-Identificado,1-Pesquisado,2-Aprovado p/ contato,3-Contatado,4-Reunião,"
              "5-Proposta/Orçamento,6-Ganho,7-Perdido,8-Em espera")
    dv = DataValidation(type="list", formula1=f'"{stages}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("H2:H500")
    dvp = DataValidation(type="list", formula1='"A-Corretores,B-Sinistros,C-Governo,D-Parceiro armazenagem"',
                         allow_blank=True)
    ws.add_data_validation(dvp)
    dvp.add("B2:B500")

    import csv
    src = Path(__file__).resolve().parent.parent / "data" / "prospects.csv"
    with src.open(encoding="utf-8") as fh:
        for i, row in enumerate(csv.DictReader(fh), start=2):
            ws.append([row["id"], row["pipeline"], row["empresa"], "", row["contato_publico"], row["fonte"],
                       row["prioridade"], "0-Identificado", "Renan", None, row["proximo_passo"], None,
                       None, None, None, row["notas"]])
    for r in range(2, 501):
        ws.cell(row=r, column=12, value=f'=IF(J{r}="","",J{r}+7)').number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=10).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=13).number_format = USD
        ws.cell(row=r, column=14).number_format = "0%"
        ws.cell(row=r, column=15, value=f'=IF(OR(M{r}="",N{r}=""),"",M{r}*N{r})').number_format = USD
    widths = [7, 16, 38, 24, 22, 45, 10, 22, 12, 13, 45, 13, 14, 12, 14, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:P500"


def build_kpis(wb):
    ws = wb.create_sheet("KPIs")
    ws["A1"] = "Painel semanal — preencher toda segunda-feira (semana anterior)."
    ws["A1"].font = Font(bold=True, size=13)
    hdr = ["Semana (segunda)", "Contatos A (corretores)", "Contatos B (sinistros)", "Ações C (governo)",
           "Reuniões", "Indicações recebidas", "Orçamentos enviados", "Contratos/jobs ganhos",
           "Receita fechada (US$)", "Mudanças realizadas", "Pack-outs realizados", "Vaults ocupados",
           "Caixa disponível (US$)", "Conversão orçamento→ganho", "Próximas 3 ações"]
    ws.append(hdr)
    style_header(ws, 2, len(hdr))
    ws.append(["META semanal (base)", 15, 10, 3, 3, 2, 4, 2, None, None, None, None, None, None, ""])
    for c in range(2, 9):
        ws.cell(row=3, column=c).fill = INPUT
    for r in range(4, 56):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=9).number_format = USD
        ws.cell(row=r, column=13).number_format = USD
        ws.cell(row=r, column=14, value=f'=IF(G{r}>0,H{r}/G{r},"")').number_format = "0%"
    ws.cell(row=57, column=1, value="Acumulado 13 semanas").font = BOLD
    for c in range(2, 12):
        col = get_column_letter(c)
        ws.cell(row=57, column=c, value=f"=SUM({col}4:{col}16)")
    ws.cell(row=57, column=9).number_format = USD
    for i in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.column_dimensions["O"].width = 60
    ws.freeze_panes = "B4"


def build_leiame(wb):
    ws = wb.create_sheet("LEIA-ME", 0)
    lines = [
        "Seckler Moving — Modelo financeiro de expansão (warehouse, storage, sinistros, governo)",
        "",
        "Como usar:",
        "1. Edite apenas as células AZUIS na aba Premissas (e metas em KPIs).",
        "2. A aba Cenarios recalcula 9 combinações: estruturas A/B/C x casos Conservador/Base/Otimista.",
        "3. A aba Sensibilidade mostra quantos vaults estáveis justificam sair do modelo terceirizado.",
        "4. CRM: altere Etapa pelo menu; Data follow-up = último contato + 7 dias (sobrescreva se precisar).",
        "5. KPIs: uma linha por semana.",
        "",
        "Regras do modelo:",
        "- Valores são mensais, em regime (após rampa). Não há projeção de crescimento automática.",
        "- Receita de governo só entra quando houver contrato assinado.",
        "- Toda premissa sem fonte está marcada como PREMISSA ou PENDENTE na coluna F.",
        "- Não é recomendação de crédito nem garantia de aprovação de financiamento.",
        "",
        "Gerado por model/build_model.py — rode novamente para recriar a planilha a partir do código.",
    ]
    for i, t in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=t)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110


def main():
    wb = Workbook()
    refs, case_hdr, struct_hdr = build_premissas(wb)
    build_cenarios(wb, refs, case_hdr, struct_hdr)
    build_sensibilidade(wb, refs, struct_hdr)
    build_crm(wb)
    build_kpis(wb)
    build_leiame(wb)
    wb.save(OUT)
    print(f"OK: {OUT}")


if __name__ == "__main__":
    main()
